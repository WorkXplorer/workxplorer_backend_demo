import logging

from django.contrib import admin
from django import forms
from django.shortcuts import redirect
from django.urls import path
from django.contrib import messages
from django.template.response import TemplateResponse

from modeltranslation.admin import TranslationAdmin
from import_export import resources
from import_export.fields import Field
from import_export.admin import ImportExportModelAdmin, ExportActionMixin

from .models import SkillCategory, Skill, SkillSynonym, LearningMaterial, SkillLearningMaterial

logger = logging.getLogger(__name__)


# ──────────────────────────────────────────────────────────
#  Standard Import/Export Resources (single-sheet operations)
# ──────────────────────────────────────────────────────────

class SkillCategoryResource(resources.ModelResource):
    """Resource for standard single-sheet SkillCategory import/export."""
    id = Field(attribute='id', column_name='ID')
    name_en = Field(attribute='name_en', column_name='name_en')
    name_uz = Field(attribute='name_uz', column_name='name_uz')
    name_ru = Field(attribute='name_ru', column_name='name_ru')
    description_en = Field(attribute='description_en', column_name='description_en')
    description_uz = Field(attribute='description_uz', column_name='description_uz')
    description_ru = Field(attribute='description_ru', column_name='description_ru')

    class Meta:
        model = SkillCategory
        fields = ('id', 'name_en', 'name_uz', 'name_ru',
                  'description_en', 'description_uz', 'description_ru')
        import_id_fields = ('name_en',)
        skip_unchanged = True
        report_skipped = True


class SkillResource(resources.ModelResource):
    """Resource for standard single-sheet Skill import/export."""
    id = Field(attribute='id', column_name='ID')
    name_en = Field(attribute='name_en', column_name='name_en')
    name_uz = Field(attribute='name_uz', column_name='name_uz')
    name_ru = Field(attribute='name_ru', column_name='name_ru')
    description_en = Field(attribute='description_en', column_name='description_en')
    description_uz = Field(attribute='description_uz', column_name='description_uz')
    description_ru = Field(attribute='description_ru', column_name='description_ru')
    category = Field(column_name='category (EN)')
    is_active = Field(attribute='is_active', column_name='is_active')

    class Meta:
        model = Skill
        fields = ('id', 'name_en', 'name_uz', 'name_ru',
                  'description_en', 'description_uz', 'description_ru', 'category', 'is_active')
        import_id_fields = ('name_en',)
        skip_unchanged = True
        report_skipped = True

    def dehydrate_category(self, skill):
        """Export: show category names as comma-separated English names."""
        return ', '.join(cat.name_en or cat.name for cat in skill.category.all())

    def after_import_row(self, row, row_result, **kwargs):
        """After importing a row, link the skill to its category."""
        if row_result.import_type in ('new', 'update'):
            category_name = row.get('category (EN)', '').strip()
            if category_name:
                try:
                    skill = Skill.objects.get(name_en=row.get('name_en', '').strip())
                    cat = SkillCategory.objects.get(name_en__iexact=category_name)
                    skill.category.add(cat)
                except (Skill.DoesNotExist, SkillCategory.DoesNotExist):
                    logger.warning(f"Could not link skill '{row.get('name_en')}' to category '{category_name}' during import.")


# ──────────────────────────────────────────────────────────
#  Multi-Sheet XLSX Upload (custom admin action)
# ──────────────────────────────────────────────────────────

class MultiSheetXlsxUploadForm(forms.Form):
    """Form for uploading multi-sheet xlsx files containing skills and categories."""
    xlsx_file = forms.FileField(
        label="XLSX File",
        help_text="Upload an .xlsx file with 'Categories' and 'Skills' sheets."
    )


def import_multi_sheet_xlsx(xlsx_file):
    """
    Parse a multi-sheet xlsx file and import categories and skills.

    Expected sheets:
      - 'Categories': columns '#', 'name_en', 'name_uz', 'name_ru'
      - 'Skills': columns '#', 'name_en', 'name_uz', 'name_ru', 'category (EN)'

    Returns a dict with import statistics.
    """
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_file, read_only=True, data_only=True)
    stats = {
        'categories_created': 0,
        'categories_updated': 0,
        'categories_skipped': 0,
        'skills_created': 0,
        'skills_updated': 0,
        'skills_skipped': 0,
        'errors': [],
    }

    # ── Step 1: Import Categories ──
    if 'Categories' in wb.sheetnames:
        ws = wb['Categories']
        rows = list(ws.iter_rows(min_row=1, values_only=True))
        if rows:
            headers = [str(h).strip().lower() if h else '' for h in rows[0]]

            name_en_idx = _find_col(headers, 'name_en')
            name_uz_idx = _find_col(headers, 'name_uz')
            name_ru_idx = _find_col(headers, 'name_ru')

            if name_en_idx is None:
                stats['errors'].append("Categories sheet: 'name_en' column not found.")
            else:
                for row_num, row in enumerate(rows[1:], start=2):
                    try:
                        name_en = _cell_str(row, name_en_idx)
                        if not name_en:
                            stats['categories_skipped'] += 1
                            continue

                        name_uz = _cell_str(row, name_uz_idx) if name_uz_idx is not None else ''
                        name_ru = _cell_str(row, name_ru_idx) if name_ru_idx is not None else ''

                        cat, created = SkillCategory.objects.update_or_create(
                            name_en__iexact=name_en,
                            defaults={
                                'name_en': name_en,
                                'name_uz': name_uz or name_en,
                                'name_ru': name_ru or name_en,
                            }
                        )
                        if created:
                            stats['categories_created'] += 1
                        else:
                            stats['categories_updated'] += 1
                    except Exception as e:
                        stats['errors'].append(f"Categories row {row_num}: {e}")
    else:
        stats['errors'].append("'Categories' sheet not found in file.")

    # ── Step 2: Import Skills ──
    if 'Skills' in wb.sheetnames:
        ws = wb['Skills']
        rows = list(ws.iter_rows(min_row=1, values_only=True))
        if rows:
            headers = [str(h).strip().lower() if h else '' for h in rows[0]]

            name_en_idx = _find_col(headers, 'name_en')
            name_uz_idx = _find_col(headers, 'name_uz')
            name_ru_idx = _find_col(headers, 'name_ru')
            cat_idx = _find_col(headers, 'category (en)')

            if name_en_idx is None:
                stats['errors'].append("Skills sheet: 'name_en' column not found.")
            else:
                # Pre-fetch all categories for efficient lookup
                cat_cache = {}
                for cat in SkillCategory.objects.all():
                    if cat.name_en:
                        cat_cache[cat.name_en.lower()] = cat

                for row_num, row in enumerate(rows[1:], start=2):
                    try:
                        name_en = _cell_str(row, name_en_idx)
                        if not name_en:
                            stats['skills_skipped'] += 1
                            continue

                        name_uz = _cell_str(row, name_uz_idx) if name_uz_idx is not None else ''
                        name_ru = _cell_str(row, name_ru_idx) if name_ru_idx is not None else ''

                        skill, created = Skill.objects.update_or_create(
                            name_en__iexact=name_en,
                            defaults={
                                'name_en': name_en,
                                'name_uz': name_uz or name_en,
                                'name_ru': name_ru or name_en,
                            }
                        )

                        # Link to category
                        if cat_idx is not None:
                            cat_name = _cell_str(row, cat_idx)
                            if cat_name:
                                cat_obj = cat_cache.get(cat_name.lower())
                                if cat_obj:
                                    skill.category.add(cat_obj)
                                else:
                                    # Auto-create category if not found
                                    new_cat, _ = SkillCategory.objects.get_or_create(
                                        name_en__iexact=cat_name,
                                        defaults={
                                            'name_en': cat_name,
                                            'name_uz': cat_name,
                                            'name_ru': cat_name,
                                        }
                                    )
                                    cat_cache[cat_name.lower()] = new_cat
                                    skill.category.add(new_cat)

                        if created:
                            stats['skills_created'] += 1
                        else:
                            stats['skills_updated'] += 1
                    except Exception as e:
                        stats['errors'].append(f"Skills row {row_num}: {e}")
    else:
        stats['errors'].append("'Skills' sheet not found in file.")

    wb.close()
    return stats


def _find_col(headers, name):
    """Find column index by header name (case-insensitive)."""
    name_lower = name.lower()
    for idx, h in enumerate(headers):
        if h == name_lower:
            return idx
    return None


def _cell_str(row, idx):
    """Safely get a cell value as stripped string."""
    if idx is not None and idx < len(row) and row[idx] is not None:
        return str(row[idx]).strip()
    return ''


# ──────────────────────────────────────────────────────────
#  Admin Classes
# ──────────────────────────────────────────────────────────

class SkillCategoryAdmin(TranslationAdmin, ImportExportModelAdmin, ExportActionMixin):
    resource_classes = [SkillCategoryResource]
    model = SkillCategory
    readonly_fields = ("created_at",)
    list_display = (
        "name",
        "description",
        "created_at",
    )
    search_fields = ("name",)


class SkillAdmin(TranslationAdmin, ImportExportModelAdmin, ExportActionMixin):
    resource_classes = [SkillResource]
    model = Skill
    readonly_fields = ("created_at",)
    show_change_link = True
    list_display = ("name", "get_categories", "description", "is_active", "created_at")
    list_filter = ("is_active",)
    search_fields = ("name",)
    change_list_template = "admin/skills/skill_changelist.html"

    def get_categories(self, obj):
        """Display all categories."""
        return ", ".join([cat.name for cat in obj.category.all()])

    get_categories.short_description = "Categories"

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path(
                'import-multi-sheet-xlsx/',
                self.admin_site.admin_view(self.import_multi_sheet_xlsx_view),
                name='skills_skill_import_multi_sheet_xlsx',
            ),
        ]
        return custom_urls + urls

    def import_multi_sheet_xlsx_view(self, request):
        """Admin view for importing multi-sheet xlsx files."""
        if request.method == 'POST':
            form = MultiSheetXlsxUploadForm(request.POST, request.FILES)
            if form.is_valid():
                xlsx_file = form.cleaned_data['xlsx_file']

                # Validate file extension
                if not xlsx_file.name.endswith('.xlsx'):
                    messages.error(request, "Only .xlsx files are supported.")
                    return redirect('..')

                try:
                    stats = import_multi_sheet_xlsx(xlsx_file)

                    # Report results
                    msg_parts = []
                    if stats['categories_created'] or stats['categories_updated']:
                        msg_parts.append(
                            f"Categories: {stats['categories_created']} created, "
                            f"{stats['categories_updated']} updated"
                        )
                    if stats['skills_created'] or stats['skills_updated']:
                        msg_parts.append(
                            f"Skills: {stats['skills_created']} created, "
                            f"{stats['skills_updated']} updated"
                        )
                    if stats['categories_skipped'] or stats['skills_skipped']:
                        msg_parts.append(
                            f"Skipped: {stats['categories_skipped']} categories, "
                            f"{stats['skills_skipped']} skills"
                        )

                    if msg_parts:
                        messages.success(request, "Import completed. " + ". ".join(msg_parts) + ".")

                    if stats['errors']:
                        for err in stats['errors'][:10]:
                            messages.warning(request, f"Import warning: {err}")
                        if len(stats['errors']) > 10:
                            messages.warning(
                                request,
                                f"... and {len(stats['errors']) - 10} more warnings."
                            )

                except Exception as e:
                    logger.exception("Multi-sheet xlsx import failed")
                    messages.error(request, f"Import failed: {e}")

                return redirect('..')
        else:
            form = MultiSheetXlsxUploadForm()

        context = {
            **self.admin_site.each_context(request),
            'form': form,
            'title': 'Import Multi-Sheet XLSX (Skills & Categories)',
            'opts': self.model._meta,
        }
        return TemplateResponse(
            request,
            'admin/skills/import_multi_sheet_xlsx.html',
            context,
        )


class SkillSynonymAdmin(admin.ModelAdmin):
    model = SkillSynonym
    show_change_link = True
    list_display = ("skill", "synonym")


class LearningMaterialAdmin(admin.ModelAdmin):
    model = LearningMaterial
    list_display = ("title", "material_type", "source", "is_free", "is_active", "rating")
    list_filter = ("material_type", "is_free", "is_active", "source")
    search_fields = ("title", "description", "url")
    readonly_fields = ("created_at", "updated_at")


class SkillLearningMaterialAdmin(admin.ModelAdmin):
    model = SkillLearningMaterial
    list_display = ("skill", "material", "relevance_score", "is_top_match")
    list_filter = ("is_top_match",)
    search_fields = ("skill__name", "material__title")
    raw_id_fields = ("skill", "material")
    readonly_fields = ("created_at", "updated_at")


admin.site.register(SkillCategory, SkillCategoryAdmin)
admin.site.register(Skill, SkillAdmin)
admin.site.register(SkillSynonym, SkillSynonymAdmin)
admin.site.register(LearningMaterial, LearningMaterialAdmin)
admin.site.register(SkillLearningMaterial, SkillLearningMaterialAdmin)
